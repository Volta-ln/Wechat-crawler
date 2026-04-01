from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import load_workbook

from app.core.extractor import CreditRecord
from app.utils.config import HEADER_ALIASES, REQUIRED_FIELDS


@dataclass
class MappingResult:
    header_row: int
    column_map: dict[str, int]


def _normalize_header(value: object) -> str:
    return str(value or "").strip()


def discover_mapping(template_path: Path) -> MappingResult:
    wb = load_workbook(template_path)
    ws = wb.active

    max_scan_rows = min(ws.max_row, 10)
    for row in range(1, max_scan_rows + 1):
        row_values = [_normalize_header(ws.cell(row=row, column=col).value) for col in range(1, ws.max_column + 1)]
        if not any(row_values):
            continue

        col_map: dict[str, int] = {}
        for idx, value in enumerate(row_values, start=1):
            if not value:
                continue
            for field, aliases in HEADER_ALIASES.items():
                if value in aliases:
                    col_map[field] = idx

        if len(col_map) >= 3:
            wb.close()
            return MappingResult(header_row=row, column_map=col_map)

    wb.close()
    raise ValueError("未识别到模板表头，请确认模板中包含中文列名")


def validate_mapping(mapping: MappingResult) -> list[str]:
    missing = [field for field in REQUIRED_FIELDS if field not in mapping.column_map]
    return missing


def _record_key_from_row(row_data: dict[str, str]) -> tuple[str, str, str, str]:
    return (
        row_data.get("推文URL", "").strip(),
        row_data.get("姓名", "").strip(),
        row_data.get("工作内容/角色", "").strip(),
        row_data.get("发布日期", "").strip(),
    )


def _record_key_from_credit(credit: CreditRecord) -> tuple[str, str, str, str]:
    return (
        credit.url.strip(),
        credit.name.strip(),
        credit.role.strip(),
        credit.publish_date.strip(),
    )


def write_records(template_path: Path, records: list[CreditRecord], output_dir: Path | None = None) -> tuple[Path, int]:
    if not records:
        raise ValueError("没有可写入的数据")

    mapping = discover_mapping(template_path)
    missing = validate_mapping(mapping)
    if missing:
        missing_text = "、".join(missing)
        raise ValueError(f"模板缺少字段: {missing_text}")

    wb = load_workbook(template_path)
    ws = wb.active

    existing_keys: set[tuple[str, str, str, str]] = set()
    start_row = mapping.header_row + 1
    for row in range(start_row, ws.max_row + 1):
        row_data: dict[str, str] = {}
        for field, col in mapping.column_map.items():
            row_data[field] = str(ws.cell(row=row, column=col).value or "").strip()
        if any(row_data.values()):
            existing_keys.add(_record_key_from_row(row_data))

    append_count = 0
    next_row = ws.max_row + 1
    for credit in records:
        key = _record_key_from_credit(credit)
        if key in existing_keys:
            continue

        ws.cell(next_row, mapping.column_map["推文标题"], credit.title)
        ws.cell(next_row, mapping.column_map["推文URL"], credit.url)
        ws.cell(next_row, mapping.column_map["发布日期"], credit.publish_date)
        ws.cell(next_row, mapping.column_map["姓名"], credit.name)
        ws.cell(next_row, mapping.column_map["工作内容/角色"], credit.role)
        ws.cell(next_row, mapping.column_map["备注"], credit.remark)

        existing_keys.add(key)
        append_count += 1
        next_row += 1

    output_base = output_dir or template_path.parent
    output_base.mkdir(parents=True, exist_ok=True)
    output_path = output_base / f"{template_path.stem}_updated_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    with NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_file:
        temp_path = Path(temp_file.name)

    wb.save(temp_path)
    wb.close()

    temp_path.replace(output_path)
    return output_path, append_count
